from pathlib import Path

from docx import Document
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models.analysis_run import AnalysisRun
from app.models.form_question_option import FormQuestionOption
from app.models.response_control_flag import ResponseControlFlag
from app.models.response_score import ResponseScore
from tests.test_analysis_orchestrator import prepare_orchestrator_fixture


def _assert_no_exe_or_u_exe() -> None:
    root = Path("E:/Colmena/backend")
    assert not list(root.rglob("u.exe"))
    assert not [path for path in root.rglob("*.exe") if ".venv" not in str(path).lower()]


def _visible_text(document: Document) -> str:
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text).strip()


def _question_option(question_id: str, label: str) -> str:
    with SessionLocal() as session:
        option = session.scalar(
            select(FormQuestionOption).where(
                FormQuestionOption.question_id == question_id,
                FormQuestionOption.label == label,
                FormQuestionOption.deleted_at.is_(None),
            )
        )
        assert option is not None
        return option.id


def test_advanced_scoring_end_to_end(client):
    fixture = prepare_orchestrator_fixture(client)
    form_id = fixture["form"]["id"]
    instrument_id = fixture["instrument"]["id"]
    dimension_id = fixture["dimension"]["id"]
    responses = fixture["responses"]
    questions = fixture["questions"]

    reverse_payload = {
        "section_id": questions["item_02"].get("section_id"),
        "instrument_id": questions["item_02"].get("instrument_id"),
        "dimension_id": questions["item_02"].get("dimension_id"),
        "project_variable_id": questions["item_02"].get("project_variable_id"),
        "code": questions["item_02"].get("code"),
        "label": questions["item_02"].get("label"),
        "help_text": questions["item_02"].get("help_text"),
        "question_type": questions["item_02"].get("question_type"),
        "question_role": questions["item_02"].get("question_role"),
        "measurement_level": questions["item_02"].get("measurement_level"),
        "data_type": questions["item_02"].get("data_type"),
        "is_required": questions["item_02"].get("is_required", False),
        "is_scored": questions["item_02"].get("is_scored", False),
        "is_reverse_scored": True,
        "min_value": questions["item_02"].get("min_value"),
        "max_value": questions["item_02"].get("max_value"),
        "sort_order": questions["item_02"].get("sort_order", 0),
        "validation_json": questions["item_02"].get("validation_json"),
        "config_json": questions["item_02"].get("config_json"),
    }
    reverse_update = client.patch(
        f"/api/v1/form-questions/{questions['item_02']['id']}",
        json=reverse_payload,
    )
    assert reverse_update.status_code == 200

    config_response = client.post(
        f"/api/v1/forms/{form_id}/scoring/configs",
        json={
            "instrument_id": instrument_id,
            "name": "Puntaje total de compromiso",
            "code": "engagement_total",
            "scoring_level": "instrument",
            "aggregation_method": "mean",
            "missing_policy": "allow_partial",
            "min_answered_items": 2,
            "min_completion_percent": 60,
            "reverse_scoring_enabled": True,
            "score_min": 1,
            "score_max": 5,
            "interpretation_enabled": True,
            "config_json": {"question_ids": [questions["item_01"]["id"], questions["item_02"]["id"], questions["item_03"]["id"]]},
            "bands": [
                {"label": "Bajo", "code": "low", "min_value": 1.0, "max_value": 2.333, "interpretation": "Nivel bajo"},
                {"label": "Medio", "code": "medium", "min_value": 2.334, "max_value": 3.666, "interpretation": "Nivel medio"},
                {"label": "Alto", "code": "high", "min_value": 3.667, "max_value": 5.0, "interpretation": "Nivel alto"},
            ],
        },
    )
    assert config_response.status_code == 201
    config = config_response.json()
    assert config["code"] == "engagement_total"

    dimension_config = client.post(
        f"/api/v1/forms/{form_id}/scoring/configs",
        json={
            "dimension_id": dimension_id,
            "name": "Puntaje de dimension",
            "code": "engagement_dimension",
            "scoring_level": "dimension",
            "aggregation_method": "mean",
            "missing_policy": "allow_partial",
            "reverse_scoring_enabled": True,
            "interpretation_enabled": True,
        },
    )
    assert dimension_config.status_code == 201

    control_scale = client.post(
        f"/api/v1/forms/{form_id}/control-scales",
        json={
            "instrument_id": instrument_id,
            "name": "Escala de mentira",
            "code": "lie_scale",
            "control_type": "lie",
            "rule_type": "count_failed",
            "threshold": 1,
            "comparison_operator": "gte",
            "flag_level": "invalid",
            "message": "Respuesta invalida por escala de control",
        },
    )
    assert control_scale.status_code == 201
    control_scale_id = control_scale.json()["id"]

    option_id = _question_option(questions["item_01"]["id"], "Likert 5")
    control_item = client.post(
        f"/api/v1/control-scales/{control_scale_id}/items",
        json={
            "question_id": questions["item_01"]["id"],
            "expected_option_id": option_id,
            "fail_if_selected": False,
            "weight": 1.0,
        },
    )
    assert control_item.status_code == 201

    preview = client.post(f"/api/v1/forms/{form_id}/scoring/preview")
    assert preview.status_code == 200
    assert preview.json()["score_results"]

    run = client.post(
        f"/api/v1/forms/{form_id}/scoring/run",
        json={"include_discarded": False, "recalculate": True, "store_result": True},
    )
    assert run.status_code == 200
    run_body = run.json()
    assert run_body["analysis_run_id"] is not None
    assert run_body["scored_responses"] == 20
    assert run_body["invalid_responses"] >= 1

    with SessionLocal() as session:
        response_score = session.scalar(
            select(ResponseScore).where(
                ResponseScore.response_id == responses[0],
                ResponseScore.scoring_config_id == config["id"],
            )
        )
        assert response_score is not None
        assert response_score.final_score == 2.0
        assert response_score.band_label == "Bajo"
        assert response_score.interpretation == "Nivel bajo"

        control_flag = session.scalar(
            select(ResponseControlFlag).where(
                ResponseControlFlag.response_id == responses[0],
                ResponseControlFlag.control_scale_id == control_scale_id,
            )
        )
        assert control_flag is not None
        assert control_flag.flag_status == "invalid"

        scoring_run = session.scalar(
            select(AnalysisRun).where(AnalysisRun.id == run_body["analysis_run_id"], AnalysisRun.analysis_type == "advanced_scoring")
        )
        assert scoring_run is not None

    response_scores = client.get(f"/api/v1/form-responses/{responses[0]}/scores")
    assert response_scores.status_code == 200
    assert any(item["scoring_config_id"] == config["id"] for item in response_scores.json())

    scoring_results = client.get(f"/api/v1/forms/{form_id}/scoring/results")
    assert scoring_results.status_code == 200
    scoring_results_body = scoring_results.json()
    assert scoring_results_body["scored_responses"] == 20
    assert scoring_results_body["band_distribution"]
    assert scoring_results_body["control_flags"]

    scoring_dataset = client.get(f"/api/v1/forms/{form_id}/scoring/dataset")
    assert scoring_dataset.status_code == 200
    dataset_body = scoring_dataset.json()
    assert "score_engagement_total" in dataset_body["columns"]
    assert "band_engagement_total" in dataset_body["columns"]

    dataset_preview = client.get(f"/api/v1/forms/{form_id}/dataset/preview?include_scores=true&limit=5")
    assert dataset_preview.status_code == 200
    preview_body = dataset_preview.json()
    assert any(column["name"] == "score_engagement_total" for column in preview_body["columns"])

    scoring_descriptives = client.get(f"/api/v1/forms/{form_id}/descriptives/scoring")
    assert scoring_descriptives.status_code == 200
    assert scoring_descriptives.json()["scored_responses"] == 20

    scoring_summary = client.post(
        f"/api/v1/forms/{form_id}/apa-tables/generate",
        json={
            "table_type": "scoring_summary",
            "source_type": "live",
            "decimals": 3,
            "include_discarded": False,
            "score_aggregation": "mean",
            "options": {},
        },
    )
    assert scoring_summary.status_code == 200
    assert scoring_summary.json()["rows"]

    scoring_from_run = client.post(f"/api/v1/forms/{form_id}/apa-tables/from-analysis-run/{run_body['analysis_run_id']}")
    assert scoring_from_run.status_code == 200
    table_types = {table["table_type"] for table in scoring_from_run.json()["tables"]}
    assert "scoring_summary" in table_types
    assert "score_band_distribution" in table_types

    scoring_chart = client.post(
        f"/api/v1/forms/{form_id}/charts/generate",
        json={
            "chart_type": "donut",
            "source_type": "live",
            "analysis_goal": "advanced_scoring",
            "theme": "colmena_premium",
            "decimals": 3,
            "include_discarded": False,
            "score_aggregation": "mean",
            "options": {"show_percentages": True},
        },
    )
    assert scoring_chart.status_code == 200
    assert scoring_chart.json()["plotly_data"]

    recommended_charts = client.get(f"/api/v1/forms/{form_id}/charts/recommended?theme=colmena_premium&max_charts=10")
    assert recommended_charts.status_code == 200
    assert any("Niveles de" in chart["title"] for chart in recommended_charts.json()["charts"])

    word_report = client.post(
        f"/api/v1/forms/{form_id}/word-reports/generate",
        json={
            "report_type": "full_form_report",
            "source_type": "mixed",
            "analysis_run_ids": [run_body["analysis_run_id"]],
            "title": "Informe con baremos",
            "subtitle": "Colmena",
            "include_charts_placeholders": True,
            "include_chart_images": False,
            "include_cover": True,
            "include_methodology_summary": True,
        },
    )
    assert word_report.status_code == 200
    word_body = word_report.json()
    docx_path = Path(word_body["file_path"])
    assert docx_path.exists()
    document = Document(docx_path)
    text = _visible_text(document)
    assert "Resultados por baremo" in text
    assert "Los niveles obtenidos corresponden a los rangos configurados" in text

    overlap_band = client.post(
        f"/api/v1/scoring/configs/{config['id']}/bands",
        json={"label": "Superpuesto", "code": "overlap", "min_value": 2.0, "max_value": 4.0, "interpretation": "Rango superpuesto"},
    )
    assert overlap_band.status_code == 201
    config_list = client.get(f"/api/v1/forms/{form_id}/scoring/configs")
    assert config_list.status_code == 200
    assert "overlapping_bands" in config_list.json()["warnings"]

    _assert_no_exe_or_u_exe()


def test_scoring_selected_images_or_no_data_safety(client):
    fixture = prepare_orchestrator_fixture(client)
    form_id = fixture["form"]["id"]

    invalid_run = client.post(
        f"/api/v1/forms/{form_id}/scoring/run",
        json={"include_discarded": False, "recalculate": True, "store_result": False},
    )
    assert invalid_run.status_code == 200
    assert invalid_run.json()["score_results"] == []
    assert invalid_run.json()["control_flags"] == []

    _assert_no_exe_or_u_exe()
